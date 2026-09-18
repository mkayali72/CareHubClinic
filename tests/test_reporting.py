"""Automated query, RBAC, boundary, soft-delete, and export coverage for Reporting."""

from datetime import date, datetime, time, timedelta
from decimal import Decimal
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import (
    Appointment,
    AppointmentStatus,
    AppointmentType,
    Clinic,
    DeliveryOutcome,
    Invoice,
    Patient,
    PregnancyEpisode,
    PregnancyEpisodeStatus,
    User,
    UserRole,
    Visit,
    VisitType,
)
from app.services.billing import create_invoice, ensure_default_fee_schedule
from app.services.clinical import create_delivery_outcome, create_pregnancy_episode
from app.services.patients import create_patient
from app.services.reporting import (
    active_pregnancies_by_trimester,
    available_reports,
    daily_schedule_census,
    delivery_outcomes_log,
    export_pdf,
    export_rows,
    export_xlsx,
    no_show_rate,
    patients_due_for_screening,
    revenue_summary,
)

REPORT_START = date(2026, 9, 12)
REPORT_END = date(2026, 9, 12)


def login_as(client: TestClient, user: User) -> None:
    """Authenticate a test client as one seeded staff user."""

    response = client.post(
        "/login",
        data={"username": user.username, "password": "Valid-Test-Password1"},
        follow_redirects=False,
    )
    assert response.status_code == 303


def make_patient(db: Session, user: User, name: str) -> Patient:
    """Create one active patient with the minimum demographic payload."""

    return create_patient(
        db=db,
        clinic_id=user.clinic_id,
        actor=user,
        payload={
            "name": name,
            "date_of_birth": date(1990, 2, 3),
            "contact_info": {},
            "insurance_info": {},
            "emergency_contact": {},
            "allergies": [],
            "current_medications": [],
            "contraception_method": None,
        },
    )


def make_visit(
    db: Session,
    user: User,
    patient: Patient,
    visit_date: date,
    *,
    gyn_data: dict | None = None,
) -> Visit:
    """Create a structured visit with optional screening due-date fields."""

    visit = Visit(
        clinic_id=user.clinic_id,
        patient_id=patient.id,
        visit_type=VisitType.GYN_ANNUAL,
        visit_date=visit_date,
        vitals={},
        prenatal_data={},
        gyn_data=gyn_data or {},
        hpi="",
        assessment="",
        plan="",
    )
    db.add(visit)
    db.flush()
    return visit


def make_appointment(
    db: Session,
    user: User,
    patient: Patient,
    appointment_type: AppointmentType,
    scheduled_at: datetime,
    status: AppointmentStatus,
) -> Appointment:
    """Create an appointment for boundary and status-mix checks."""

    appointment = Appointment(
        clinic_id=user.clinic_id,
        patient_id=patient.id,
        doctor_id=user.id,
        scheduled_at=scheduled_at,
        duration_minutes=30,
        appointment_type_id=appointment_type.id,
        status=status,
    )
    db.add(appointment)
    db.flush()
    return appointment


def enable_billing(db: Session, user: User) -> None:
    """Enable Billing for financial-report setup."""

    clinic = db.scalar(select(Clinic).where(Clinic.id == user.clinic_id))
    assert clinic is not None
    clinic.billing_module_enabled = True
    db.commit()


def test_schedule_and_no_show_use_inclusive_boundaries_and_exclude_deleted_rows(
    seeded_users: dict[UserRole, User],
    db_session: Session,
) -> None:
    """End-of-day appointments count; cancelled/deleted records do not skew rate."""

    admin = seeded_users[UserRole.CLINIC_ADMIN]
    patient = make_patient(db_session, admin, "Schedule Patient")
    appointment_type = AppointmentType(
        clinic_id=admin.clinic_id,
        name="Report visit",
        default_duration_minutes=30,
    )
    db_session.add(appointment_type)
    db_session.flush()
    make_appointment(
        db_session,
        admin,
        patient,
        appointment_type,
        datetime.combine(REPORT_END, time(23, 59)),
        AppointmentStatus.NO_SHOW,
    )
    make_appointment(
        db_session,
        admin,
        patient,
        appointment_type,
        datetime.combine(REPORT_START, time(9, 0)),
        AppointmentStatus.CANCELLED,
    )
    deleted = make_appointment(
        db_session,
        admin,
        patient,
        appointment_type,
        datetime.combine(REPORT_START, time(10, 0)),
        AppointmentStatus.DONE,
    )
    deleted.soft_delete(admin.id)
    db_session.commit()

    schedule = daily_schedule_census(
        db_session,
        admin,
        REPORT_START,
        REPORT_END,
    )
    no_show = no_show_rate(db_session, admin, REPORT_START, REPORT_END)
    assert schedule["total_appointments"] == 2
    assert schedule["days"][0]["status_counts"]["no_show"] == 1
    assert schedule["days"][0]["status_counts"]["cancelled"] == 1
    assert no_show["total_appointments"] == 2
    assert no_show["cancelled_count"] == 1
    assert no_show["eligible_count"] == 1
    assert no_show["no_show_count"] == 1
    assert no_show["rate_percent"] == Decimal("100.00")


def test_reports_are_role_filtered_and_revenue_requires_enabled_billing(
    client: TestClient,
    seeded_users: dict[UserRole, User],
    db_session: Session,
) -> None:
    """Financial reports are restricted, while clinical reports stay private."""

    admin = seeded_users[UserRole.CLINIC_ADMIN]
    assert {report["slug"] for report in available_reports(db_session, admin)} == {
        "schedule",
        "no-show-rate",
        "screening",
        "pregnancies",
        "deliveries",
    }
    login_as(client, seeded_users[UserRole.PHYSICIAN])
    assert client.get("/reports/revenue").status_code == 403
    assert client.get("/reports/screening").status_code == 200
    login_as(client, seeded_users[UserRole.BILLING_CLERK])
    assert client.get("/reports/screening").status_code == 403
    assert client.get("/reports/revenue").status_code == 404

    clinic = db_session.scalar(select(Clinic).where(Clinic.id == admin.clinic_id))
    assert clinic is not None
    clinic.billing_module_enabled = True
    db_session.commit()
    assert client.get("/reports/revenue").status_code == 200
    xlsx = client.get(
        "/reports/revenue/export/xlsx?start_date=2026-09-12&end_date=2026-09-12"
    )
    pdf = client.get(
        "/reports/revenue/export/pdf?start_date=2026-09-12&end_date=2026-09-12"
    )
    assert xlsx.status_code == 200
    assert xlsx.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert pdf.status_code == 200
    assert pdf.headers["content-type"] == "application/pdf"


def test_screening_pregnancy_and_delivery_reports_are_hand_countable(
    seeded_users: dict[UserRole, User],
    db_session: Session,
) -> None:
    """Clinical reports use the documented date and trimester rules."""

    physician = seeded_users[UserRole.PHYSICIAN]
    patient = make_patient(db_session, physician, "Clinical Report Patient")
    make_visit(
        db_session,
        physician,
        patient,
        REPORT_END,
        gyn_data={
            "pap_due_date": "2026-09-12",
            "hpv_due_date": "2026-12-01",
        },
    )
    episode = create_pregnancy_episode(
        db_session,
        physician,
        patient.id,
        lmp=date(2026, 6, 1),
        edd=date(2027, 3, 8),
    )
    db_session.flush()
    create_delivery_outcome(
        db_session,
        episode,
        physician,
        date(2026, 9, 12),
        "vaginal",
        "None",
        3200,
        8,
        9,
    )
    create_pregnancy_episode(
        db_session,
        physician,
        patient.id,
        lmp=date(2026, 6, 1),
        edd=date(2027, 3, 8),
    )
    db_session.commit()

    screening = patients_due_for_screening(
        db_session,
        physician,
        REPORT_START,
        REPORT_END,
    )
    pregnancies = active_pregnancies_by_trimester(
        db_session,
        physician,
        REPORT_START,
        REPORT_END,
    )
    deliveries = delivery_outcomes_log(
        db_session,
        physician,
        REPORT_START,
        REPORT_END,
    )
    assert screening["patient_count"] == 1
    assert screening["rows"][0]["overdue"] == "Pap"
    assert pregnancies["total"] == 1
    assert pregnancies["counts"]["second"] == 1
    assert pregnancies["rows"][0]["gestational_age"] == "14w 5d"
    assert deliveries["count"] == 1
    assert deliveries["rows"][0]["mode"] == "vaginal"


def test_revenue_is_inclusive_and_excludes_soft_deleted_financial_records(
    seeded_users: dict[UserRole, User],
    db_session: Session,
) -> None:
    """Revenue totals use invoice snapshots and omit a deleted invoice."""

    admin = seeded_users[UserRole.CLINIC_ADMIN]
    enable_billing(db_session, admin)
    patient = make_patient(db_session, admin, "Revenue Patient")
    visit = make_visit(db_session, admin, patient, REPORT_END)
    fee = ensure_default_fee_schedule(db_session, admin.clinic_id)[0]
    db_session.commit()
    invoice = create_invoice(db_session, admin, visit.id, [fee.id])
    invoice.created_at = datetime.combine(REPORT_END, time(23, 59))
    db_session.commit()

    report = revenue_summary(db_session, admin, REPORT_START, REPORT_END)
    assert report["invoice_count"] == 1
    assert report["gross_total"] == Decimal("150.00")
    invoice.soft_delete(admin.id)
    db_session.commit()
    assert revenue_summary(
        db_session,
        admin,
        REPORT_START,
        REPORT_END,
    )["invoice_count"] == 0


def test_pdf_and_excel_exports_are_real_downloads(
    seeded_users: dict[UserRole, User],
    db_session: Session,
) -> None:
    """Export helpers produce parseable XLSX and a valid PDF signature."""

    admin = seeded_users[UserRole.CLINIC_ADMIN]
    patient = make_patient(db_session, admin, "Export Patient")
    make_visit(
        db_session,
        admin,
        patient,
        REPORT_END,
        gyn_data={"pap_due_date": "2026-09-01"},
    )
    db_session.commit()
    report = patients_due_for_screening(
        db_session,
        admin,
        REPORT_START,
        REPORT_END,
    )
    workbook_response = export_xlsx(report)
    workbook = load_workbook(filename=__import__("io").BytesIO(workbook_response.body))
    assert workbook.active["A1"].value == "Patients due for screening"
    assert workbook.active["A4"].value == "Patient ID"
    pdf_response = export_pdf(report)
    assert pdf_response.body.startswith(b"%PDF-")
    assert pdf_response.headers["content-disposition"].endswith('"screening.pdf"')


def test_92_known_dataset_returns_exact_values_for_every_report(
    seeded_users: dict[UserRole, User],
    db_session: Session,
) -> None:
    """Test 92: every report is checked against a hand-counted dataset."""

    admin = seeded_users[UserRole.CLINIC_ADMIN]
    enable_billing(db_session, admin)
    patients = [
        make_patient(db_session, admin, "Known Appointment Patient 1"),
        make_patient(db_session, admin, "Known Appointment Patient 2"),
        make_patient(db_session, admin, "Known Appointment Patient 3"),
    ]
    appointment_type = AppointmentType(
        clinic_id=admin.clinic_id,
        name="Known report appointment",
        default_duration_minutes=30,
    )
    db_session.add(appointment_type)
    db_session.flush()
    statuses = (
        AppointmentStatus.NO_SHOW,
        AppointmentStatus.NO_SHOW,
        AppointmentStatus.DONE,
    )
    for index, (patient, status) in enumerate(zip(patients, statuses)):
        make_appointment(
            db_session,
            admin,
            patient,
            appointment_type,
            datetime.combine(REPORT_END, time(9 + index)),
            status,
        )
    screening_visit = make_visit(
        db_session,
        admin,
        patients[0],
        REPORT_END,
        gyn_data={
            "pap_due_date": REPORT_END.isoformat(),
            "hpv_due_date": "2026-12-01",
        },
    )
    fee = ensure_default_fee_schedule(db_session, admin.clinic_id)[0]
    db_session.commit()
    invoice = create_invoice(db_session, admin, screening_visit.id, [fee.id])
    invoice.created_at = datetime.combine(REPORT_END, time(23, 59))
    delivery_episode = create_pregnancy_episode(
        db_session,
        admin,
        patients[1].id,
        lmp=date(2025, 12, 1),
        edd=date(2026, 9, 7),
    )
    create_delivery_outcome(
        db_session,
        delivery_episode,
        admin,
        REPORT_END,
        "vaginal",
        "None",
        3200,
        8,
        9,
    )
    create_pregnancy_episode(
        db_session,
        admin,
        patients[2].id,
        lmp=date(2026, 6, 1),
        edd=date(2027, 3, 8),
    )
    db_session.commit()

    schedule = daily_schedule_census(db_session, admin, REPORT_START, REPORT_END)
    revenue = revenue_summary(db_session, admin, REPORT_START, REPORT_END)
    no_show = no_show_rate(db_session, admin, REPORT_START, REPORT_END)
    screening = patients_due_for_screening(db_session, admin, REPORT_START, REPORT_END)
    pregnancies = active_pregnancies_by_trimester(
        db_session,
        admin,
        REPORT_START,
        REPORT_END,
    )
    deliveries = delivery_outcomes_log(db_session, admin, REPORT_START, REPORT_END)

    assert schedule["total_appointments"] == 3
    assert schedule["days"][0]["total_appointments"] == 3
    assert schedule["days"][0]["active_census"] == 1
    assert schedule["days"][0]["status_counts"] == {
        "scheduled": 0,
        "checked_in": 0,
        "in_room": 0,
        "with_doctor": 0,
        "done": 1,
        "cancelled": 0,
        "no_show": 2,
    }
    assert revenue["invoice_count"] == 1
    assert revenue["gross_total"] == Decimal("150.00")
    assert revenue["paid_total"] == Decimal("0.00")
    assert revenue["unpaid_total"] == Decimal("150.00")
    assert no_show["eligible_count"] == 3
    assert no_show["no_show_count"] == 2
    assert no_show["rate_percent"] == Decimal("66.67")
    assert screening["patient_count"] == 1
    assert screening["rows"][0]["patient"] == "Known Appointment Patient 1"
    assert screening["rows"][0]["overdue"] == "Pap"
    assert pregnancies["total"] == 1
    assert pregnancies["counts"] == {"first": 0, "second": 1, "third": 0}
    assert pregnancies["rows"][0]["patient"] == "Known Appointment Patient 3"
    assert deliveries["count"] == 1
    assert deliveries["rows"][0]["patient"] == "Known Appointment Patient 2"
    assert deliveries["rows"][0]["mode"] == "vaginal"


def test_93_date_filter_includes_both_boundaries_and_excludes_neighbors(
    seeded_users: dict[UserRole, User],
    db_session: Session,
) -> None:
    """Test 93: inclusive date boundaries are asserted with adjacent records."""

    admin = seeded_users[UserRole.CLINIC_ADMIN]
    patient = make_patient(db_session, admin, "Boundary Patient")
    appointment_type = AppointmentType(
        clinic_id=admin.clinic_id,
        name="Boundary appointment",
        default_duration_minutes=30,
    )
    db_session.add(appointment_type)
    db_session.flush()
    inside_start = make_appointment(
        db_session,
        admin,
        patient,
        appointment_type,
        datetime.combine(REPORT_START, time.min),
        AppointmentStatus.NO_SHOW,
    )
    inside_end = make_appointment(
        db_session,
        admin,
        patient,
        appointment_type,
        datetime.combine(REPORT_END, time(23, 59, 59)),
        AppointmentStatus.NO_SHOW,
    )
    outside_before = make_appointment(
        db_session,
        admin,
        patient,
        appointment_type,
        datetime.combine(REPORT_START, time.min) - timedelta(seconds=1),
        AppointmentStatus.DONE,
    )
    outside_after = make_appointment(
        db_session,
        admin,
        patient,
        appointment_type,
        datetime.combine(REPORT_END, time.min) + timedelta(days=1),
        AppointmentStatus.DONE,
    )
    db_session.commit()

    schedule = daily_schedule_census(db_session, admin, REPORT_START, REPORT_END)
    no_show = no_show_rate(db_session, admin, REPORT_START, REPORT_END)
    assert {row["id"] for row in schedule["days"][0]["rows"]} == {
        inside_start.id,
        inside_end.id,
    }
    assert outside_before.id not in {
        row["id"] for row in schedule["days"][0]["rows"]
    }
    assert outside_after.id not in {
        row["id"] for row in schedule["days"][0]["rows"]
    }
    assert no_show["total_appointments"] == 2
    assert no_show["no_show_count"] == 2
    assert no_show["rate_percent"] == Decimal("100.00")


def test_94_pdf_and_excel_data_match_the_same_report_rows(
    seeded_users: dict[UserRole, User],
    db_session: Session,
) -> None:
    """Test 94: parsed Excel and PDF bytes contain the displayed report data."""

    admin = seeded_users[UserRole.CLINIC_ADMIN]
    patient = make_patient(db_session, admin, "Export Match Patient")
    appointment_type = AppointmentType(
        clinic_id=admin.clinic_id,
        name="Export match appointment",
        default_duration_minutes=30,
    )
    db_session.add(appointment_type)
    db_session.flush()
    make_appointment(
        db_session,
        admin,
        patient,
        appointment_type,
        datetime.combine(REPORT_END, time(23, 59)),
        AppointmentStatus.NO_SHOW,
    )
    db_session.commit()
    report = daily_schedule_census(db_session, admin, REPORT_START, REPORT_END)
    columns, expected_rows = export_rows(report)

    excel = export_xlsx(report)
    workbook = load_workbook(filename=BytesIO(excel.body), data_only=True)
    worksheet = workbook.active
    assert list(next(worksheet.iter_rows(min_row=4, max_row=4, values_only=True))) == columns
    parsed_rows = [
        list(row)
        for row in worksheet.iter_rows(min_row=5, values_only=True)
    ]
    normalized_rows = [
        [
            value.date()
            if isinstance(value, datetime) and value.time() == time.min
            else value
            for value in row
        ]
        for row in parsed_rows
    ]
    assert normalized_rows == expected_rows

    pdf = export_pdf(report)
    pdf_text = "\n".join(
        page.extract_text() or "" for page in PdfReader(BytesIO(pdf.body)).pages
    )
    assert report["title"] in pdf_text
    for value in expected_rows[0]:
        assert str(value) in pdf_text


def test_95_front_desk_cannot_read_financial_report_data_directly(
    client: TestClient,
    seeded_users: dict[UserRole, User],
    db_session: Session,
) -> None:
    """Test 95: direct financial report and export calls enforce RBAC."""

    admin = seeded_users[UserRole.CLINIC_ADMIN]
    enable_billing(db_session, admin)
    login_as(client, seeded_users[UserRole.FRONT_DESK])
    report_response = client.get(
        "/reports/revenue?start_date=2026-09-12&end_date=2026-09-12"
    )
    export_response = client.get(
        "/reports/revenue/export/xlsx?start_date=2026-09-12&end_date=2026-09-12"
    )
    assert report_response.status_code == 403
    assert export_response.status_code == 403


def test_96_soft_deleted_appointment_is_removed_from_report_counts(
    seeded_users: dict[UserRole, User],
    db_session: Session,
) -> None:
    """Test 96: soft deletion removes a counted appointment without hard delete."""

    admin = seeded_users[UserRole.CLINIC_ADMIN]
    patient = make_patient(db_session, admin, "Soft Delete Report Patient")
    appointment_type = AppointmentType(
        clinic_id=admin.clinic_id,
        name="Soft delete report appointment",
        default_duration_minutes=30,
    )
    db_session.add(appointment_type)
    db_session.flush()
    retained = make_appointment(
        db_session,
        admin,
        patient,
        appointment_type,
        datetime.combine(REPORT_END, time(9)),
        AppointmentStatus.DONE,
    )
    removed = make_appointment(
        db_session,
        admin,
        patient,
        appointment_type,
        datetime.combine(REPORT_END, time(10)),
        AppointmentStatus.DONE,
    )
    db_session.commit()
    assert daily_schedule_census(
        db_session,
        admin,
        REPORT_START,
        REPORT_END,
    )["total_appointments"] == 2

    removed.soft_delete(admin.id)
    db_session.commit()
    report = daily_schedule_census(
        db_session,
        admin,
        REPORT_START,
        REPORT_END,
    )
    assert report["total_appointments"] == 1
    assert report["days"][0]["rows"][0]["id"] == retained.id
    assert db_session.get(Appointment, removed.id).deleted_at is not None